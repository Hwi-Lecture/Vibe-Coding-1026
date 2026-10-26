"""Build the four-sheet workbook from saved CSV files, without API calls.

Install: python -m pip install openpyxl
Run: python make_opinet_excel.py --as-of 2026-10-23
Without --as-of, the most recent dated data folder is used.
"""

import argparse
import csv
import json
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
PRODUCTS = {'B027': '보통휘발유', 'B034': '고급휘발유', 'D047': '자동차용경유',
            'C004': '실내등유', 'K015': '자동차용부탄'}
BRANDS = {'SKE': 'SK에너지', 'GSC': 'GS칼텍스', 'HDO': 'HD현대오일뱅크',
          'SOL': 'S-OIL', 'RTO': '알뜰주유소', 'RTE': '자영알뜰',
          'RTX': '알뜰(ex)', 'NHO': 'NH-OIL', 'SKG': 'SK가스', 'E1G': 'E1', 'ETC': '자가상표'}
NAVY = '203864'
PRICE_FORMAT = '#,##0.00'


def read_data(folder, as_of):
    datasets = defaultdict(list)
    seen = set()
    manifest_path = folder / 'manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig')) if manifest_path.is_file() else {}
    entries = manifest.get('files', [])
    if manifest and manifest.get('status') != 'complete':
        raise ValueError('수집이 완료된 자료만 엑셀로 만들 수 있습니다.')
    if entries and all(entry['path'].endswith('.csv') for entry in entries):
        filenames = [folder / entry['path'] for entry in entries]
    else:
        csv_folder = folder / 'csv' if (folder / 'csv').is_dir() else folder
        filenames = sorted(csv_folder.rglob('*.csv'))
    for filename in filenames:
        with filename.open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                if 'DATE' not in row:  # Region-code response is a lookup, not a price record.
                    continue
                day = datetime.strptime(row['DATE'], '%Y%m%d').date()
                age = (as_of - day).days
                if not 1 <= age <= 14:
                    raise ValueError(f'Date outside comparison periods: {filename}, {day}')
                period = '최근 7일' if age <= 7 else '직전 7일'
                product = row['PRODCD']
                kind = '지역별 일별' if 'AREA_CD' in row else '상표별 일별' if 'POLL_DIV_CD' in row else '전국 일별'
                category = row.get('AREA_CD', row.get('POLL_DIV_CD', ''))
                key = kind, day, product, category
                if key in seen:
                    raise ValueError(f'Duplicate source record: {key}')
                seen.add(key)
                values = [period, day, PRODUCTS.get(product, product), float(row['PRICE']), product]
                if kind == '지역별 일별':
                    values += [row['AREA_NM'], category]
                elif kind == '상표별 일별':
                    values += [BRANDS.get(category, category), category]
                datasets[kind].append(values)
    for kind in ['전국 일별', '상표별 일별', '지역별 일별']:
        if not datasets[kind]:
            raise ValueError(f'Missing data: {kind}')
        datasets[kind].sort(key=lambda row: (row[1], row[4], row[6] if len(row) > 5 else ''))
    return datasets


def header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='맑은 고딕', size=10, color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def write_sources(wb, datasets):
    for index, kind in enumerate(['전국 일별', '상표별 일별', '지역별 일별'], 1):
        rows = datasets[kind]
        ws = wb.create_sheet(kind)
        ws.sheet_view.showGridLines = False
        ws['A1'] = '출처: 한국석유공사 오피넷 API / 저장된 CSV / 가격 단위: 원/L'
        labels = ['기간 구분', '날짜', '제품', '일평균 가격(원/L)', '제품코드']
        if kind == '지역별 일별':
            labels += ['지역', '지역코드']
        elif kind == '상표별 일별':
            labels += ['상표', '상표코드']
        header(ws, 3, labels)
        for values in rows:
            ws.append(values)
        for cells in ws.iter_rows(min_row=4):
            for cell in cells:
                cell.font = Font(name='맑은 고딕', size=10)
                cell.alignment = Alignment(vertical='center')
            cells[1].number_format = 'yyyy-mm-dd'
            cells[1].alignment = Alignment(horizontal='center', vertical='center')
            cells[3].number_format = PRICE_FORMAT
            cells[4].number_format = '@'
            if len(cells) > 5:
                cells[6].number_format = '@'
            ws.row_dimensions[cells[0].row].height = 20
        for col, width in {'A': 14, 'B': 15, 'C': 20, 'D': 23, 'E': 14, 'F': 24, 'G': 14}.items():
            ws.column_dimensions[col].width = width
        table = Table(displayName=f'PriceData{index}', ref=f'A3:{chr(64 + len(labels))}{ws.max_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = 'C4'


def make_summary(wb, datasets, as_of):
    ws = wb.active
    ws.title = '기간별 비교'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws['A2'] = '유가 기간별 비교'
    ws['A2'].font = Font(name='맑은 고딕', size=16, bold=True, color=NAVY)
    ws['A3'] = f"직전: {as_of - timedelta(days=14):%Y-%m-%d} ~ {as_of - timedelta(days=8):%Y-%m-%d}"
    ws['D3'] = f"최근: {as_of - timedelta(days=7):%Y-%m-%d} ~ {as_of - timedelta(days=1):%Y-%m-%d}"
    ws['A4'] = '7개 일평균 가격의 단순평균입니다. 공식 주간 통계와 다를 수 있습니다.'
    ws['A5'] = '관측일이 7일 미만이면 평균을 표시하지 않습니다. 0일은 해당 제품 자료가 없는 경우입니다.'
    ws['A7'] = '전국 제품별'
    labels = ['제품', '직전 평균(원/L)', '최근 평균(원/L)', '증감액(원/L)', '증감률', '직전 관측일', '최근 관측일']
    header(ws, 8, labels)
    ws['A15'] = '비교할 제품'
    ws['B15'] = '보통휘발유'
    ws['B15'].fill = PatternFill('solid', fgColor='FFF2CC')
    selector = DataValidation(type='list', formula1='"' + ','.join(PRODUCTS.values()) + '"')
    selector.errorTitle = '제품 선택'
    selector.error = '목록에 있는 제품을 선택하세요.'
    selector.showErrorMessage = True
    selector.errorStyle = 'stop'
    ws.add_data_validation(selector)
    selector.add(ws['B15'])

    def section(source, start, items, category=False):
        last = 3 + len(datasets[source])
        prefix = f"'{source}'!"
        prices = f'{prefix}$D$4:$D${last}'
        periods = f'{prefix}$A$4:$A${last}'
        products = f'{prefix}$C$4:$C${last}'
        categories = f'{prefix}$F$4:$F${last}'
        for r, item in enumerate(items, start):
            ws.cell(r, 1, item)
            product = '$B$15' if category else f'$A{r}'
            extra = f',{categories},$A{r}' if category else ''
            for avg_col, count_col, period in [('B', 'F', '직전 7일'), ('C', 'G', '최근 7일')]:
                criteria = f'{periods},"{period}",{products},{product}{extra}'
                ws[f'{count_col}{r}'] = f'=COUNTIFS({criteria})'
                ws[f'{avg_col}{r}'] = f'=IF({count_col}{r}=7,AVERAGEIFS({prices},{criteria}),"")'
                ws[f'{avg_col}{r}'].number_format = PRICE_FORMAT
            ws[f'D{r}'] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),C{r}-B{r},"")'
            ws[f'E{r}'] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(B{r})),IF(B{r}=0,"",D{r}/B{r}),"")'
            ws[f'D{r}'].number_format = '+#,##0.00;-#,##0.00;0.00'
            ws[f'E{r}'].number_format = '+0.00%;-0.00%;0.00%'
        end = start + len(items) - 1
        for operator, color in [('greaterThan', 'C00000'), ('lessThan', '0070C0')]:
            ws.conditional_formatting.add(f'D{start}:E{end}', CellIsRule(
                operator=operator, formula=['0'], font=Font(color=color)))

    section('전국 일별', 9, list(PRODUCTS.values()))
    brands = sorted({row[5] for row in datasets['상표별 일별']})
    ws['A17'] = '전국 제품·상표별'
    header(ws, 18, ['상표'] + labels[1:])
    section('상표별 일별', 19, brands, True)
    region_title = 19 + len(brands) + 2
    ws.cell(region_title, 1, '제품·지역별')
    header(ws, region_title + 1, ['지역'] + labels[1:])
    regions = sorted({row[5] for row in datasets['지역별 일별']})
    section('지역별 일별', region_title + 2, regions, True)
    ws.cell(ws.max_row + 2, 1, '전국 평균은 전국 데이터를 사용합니다. 지역별 평균을 다시 평균 내지 않습니다.')
    for row in ws.iter_rows():
        for cell in row:
            if cell.row != 2 and cell.fill.fgColor.rgb != '00203864':
                cell.font = Font(name='맑은 고딕', size=10)
            is_header = cell.row in [8, 18, region_title + 1]
            cell.alignment = Alignment(vertical='center', wrap_text=is_header,
                                       horizontal='center' if is_header else 'right' if cell.data_type == 'f' else 'left')
        if row[0].row not in [8, 18, region_title + 1]:
            ws.row_dimensions[row[0].row].height = 24
    for col, width in {'A': 26, 'B': 22, 'C': 22, 'D': 23, 'E': 17, 'F': 16, 'G': 16}.items():
        ws.column_dimensions[col].width = width
    ws.sheet_view.zoomScale = 85


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--as-of', type=date.fromisoformat)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    data_root = ROOT / 'data' / 'opinet'
    as_of = args.as_of or max(date.fromisoformat(p.name) for p in data_root.iterdir()
                             if p.is_dir() and len(p.name) == 10)
    folder = data_root / as_of.isoformat()
    datasets = read_data(folder, as_of)
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    write_sources(wb, datasets)
    make_summary(wb, datasets, as_of)
    destination = args.output or folder / '유가_기간비교.xlsx'
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    print(f'Saved workbook: {destination}')
    print('Daily records: ' + ', '.join(f'{name}={len(rows)}' for name, rows in datasets.items()))


if __name__ == '__main__':
    main()
