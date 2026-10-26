"""Collect CSV data, build an Excel workbook, and send it using Gmail SMTP.

Install: python -m pip install openpyxl
Run all steps: python run_opinet.py --to recipient@example.com
Choose date: python run_opinet.py --to recipient@example.com --as-of 2026-10-23
Use saved CSV without sending: python run_opinet.py --to recipient@example.com --as-of 2026-10-23 --skip-collect --dry-run

.env (next to this file): opinet_api_key, email_address, gmail_app_password.
--dry-run skips SMTP only. Add --skip-collect to also skip API requests.
This file is standalone; the three original Python files are not required.
"""

import argparse
import csv
import json
import os
import smtplib
import ssl
import sys
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from email.headerregistry import Address
from email.message import EmailMessage
from email.utils import formatdate, make_msgid
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
KST = timezone(timedelta(hours=9))

PRODUCTS = {'B027': '보통휘발유', 'B034': '고급휘발유', 'D047': '자동차용경유',
            'C004': '실내등유', 'K015': '자동차용부탄'}
BRANDS = {'SKE': 'SK에너지', 'GSC': 'GS칼텍스', 'HDO': 'HD현대오일뱅크',
          'SOL': 'S-OIL', 'RTO': '알뜰주유소', 'RTE': '자영알뜰',
          'RTX': '알뜰(ex)', 'NHO': 'NH-OIL', 'SKG': 'SK가스', 'E1G': 'E1', 'ETC': '자가상표'}
NAVY = '203864'
PRICE_FORMAT = '#,##0.00'



def load_key():
    values = {}
    project_root = ROOT.parent.parent if ROOT.parent.name == 'checkpoint' else ROOT.parent
    for env_path in dict.fromkeys([project_root / '.env', ROOT / '.env']):
        if env_path.is_file():
            for line in env_path.read_text(encoding='utf-8-sig').splitlines():
                if '=' in line and not line.lstrip().startswith('#'):
                    name, value = line.split('=', 1)
                    values[name.strip().lower()] = value.strip().strip(chr(34)).strip(chr(39))
    key = os.environ.get('OPINET_API_KEY') or values.get('opinet_api_key')
    if not key:
        raise RuntimeError('Set opinet_api_key in .env or OPINET_API_KEY in the environment.')
    return key


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    temporary.replace(path)


def save_csv(path, rows):
    """Keep original columns and values; use a BOM for Korean text in Excel."""
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = list(dict.fromkeys(key for row in rows for key in row))
    temporary = path.with_suffix('.tmp')
    with temporary.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def load_settings():
    settings = {}
    project_root = ROOT.parent.parent if ROOT.parent.name == 'checkpoint' else ROOT.parent
    for env_file in dict.fromkeys([project_root / '.env', ROOT / '.env']):
        if env_file.is_file():
            for line in env_file.read_text(encoding='utf-8-sig').splitlines():
                if '=' in line and not line.lstrip().startswith('#'):
                    key, value = line.split('=', 1)
                    settings[key.strip().lower()] = value.strip().strip(chr(34)).strip(chr(39))
    sender = os.environ.get('EMAIL_ADDRESS') or settings.get('email_address', '')
    password = os.environ.get('GMAIL_APP_PASSWORD') or settings.get('gmail_app_password', '')
    # Google displays app passwords in groups separated by spaces.
    password = ''.join(password.split())
    if not sender or not password:
        raise ValueError('.env에 email_address와 gmail_app_password를 설정하세요.')
    validate_address(sender)
    return sender, password


def validate_address(value):
    if '\r' in value or '\n' in value:
        raise ValueError('이메일 주소에 줄바꿈을 사용할 수 없습니다.')
    try:
        parsed = Address(addr_spec=value)
        if not parsed.username or not parsed.domain:
            raise ValueError
        value.encode('ascii')
    except (ValueError, IndexError, UnicodeError):
        raise ValueError('이메일 주소를 name@example.com 형식으로 입력하세요.') from None


def build_message(sender, recipients, attachment, subject=None, body=None):
    if not attachment.is_file():
        raise FileNotFoundError('첨부할 엑셀 파일을 찾을 수 없습니다.')
    if attachment.suffix.lower() != '.xlsx' or not zipfile.is_zipfile(attachment):
        raise ValueError('첨부파일은 정상적인 .xlsx 파일이어야 합니다.')
    with zipfile.ZipFile(attachment) as archive:
        if 'xl/workbook.xml' not in archive.namelist():
            raise ValueError('첨부파일에 엑셀 통합문서가 없습니다.')
    message = EmailMessage()
    message['From'] = sender
    message['To'] = ', '.join(recipients)
    message['Subject'] = subject if subject is not None else '오피넷 유가 기간별 비교 자료'
    message['Date'] = formatdate(localtime=True)
    message['Message-ID'] = make_msgid()
    message.set_content(body if body is not None else (
        '안녕하세요.\n\n오피넷 유가 기간별 비교 엑셀 파일을 첨부합니다.\n'
        '파일에는 기간별 비교, 전국 일별, 상표별 일별, 지역별 일별 자료가 담겨 있습니다.\n\n감사합니다.\n'))
    message.add_attachment(
        attachment.read_bytes(),
        maintype='application',
        subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename=attachment.name,
    )
    return message


def collect_data(as_of, key, output):
    manifest = {'as_of': as_of.isoformat(),
                'collected_at': datetime.now(KST).isoformat(),
                'status': 'collecting', 'request_count': 0, 'files': [], 'warnings': []}

    def fetch(endpoint, filename, params=None, expected_dates=None):
        params = params or {}
        manifest['request_count'] += 1
        save_json(output / 'manifest.json', manifest)
        url = 'https://www.opinet.co.kr/api/' + endpoint + '?' + urlencode(
            {'code': key, 'out': 'json', **params})
        try:
            with urlopen(url, timeout=30) as response:
                payload = json.loads(response.read().decode('utf-8-sig'))
        except Exception as exc:
            # Exception messages can contain the request URL and API key.
            raise RuntimeError(f'{endpoint}: request failed ({type(exc).__name__})') from None
        rows = payload.get('RESULT', {}).get('OIL')
        if not isinstance(rows, list) or not rows:
            raise RuntimeError(f'{endpoint}: no data returned; stopping collection.')
        if key in json.dumps(payload, ensure_ascii=False):
            raise RuntimeError(f'{endpoint}: response unexpectedly contains the API key.')
        if expected_dates is not None:
            actual = {str(row.get('DATE')) for row in rows}
            if actual != expected_dates:
                raise RuntimeError(f'{endpoint}: unexpected date coverage: {sorted(actual)}')
            groups = {}
            for row in rows:
                group = (row.get('PRODCD'), row.get('POLL_DIV_CD'), row.get('AREA_CD'))
                group_dates = groups.setdefault(group, set())
                if row['DATE'] in group_dates:
                    raise RuntimeError(f'{endpoint}: duplicate daily record in {group}')
                group_dates.add(row['DATE'])
            for group, group_dates in groups.items():
                if group_dates != expected_dates:
                    manifest['warnings'].append({'file': filename, 'group': group,
                                                 'missing_dates': sorted(expected_dates - group_dates)})
        save_csv(output / filename, rows)
        manifest['files'].append({'path': filename, 'endpoint': endpoint,
                                  'parameters': params, 'row_count': len(rows)})
        save_json(output / 'manifest.json', manifest)
        print(f'{filename}: {len(rows)} records', flush=True)
        return rows

    try:
        areas = fetch('areaCode.do', 'area_codes.csv')
        codes = [str(row['AREA_CD']) for row in areas]
        if len(set(codes)) != len(codes) or any(len(code) != 2 or not code.isdigit() for code in codes):
            raise RuntimeError('Invalid province code list.')
        for period, days_back in [('recent', 1), ('previous', 8)]:
            end = as_of - timedelta(days=days_back)
            start = end - timedelta(days=6)
            dates = {(start + timedelta(days=n)).strftime('%Y%m%d') for n in range(7)}
            params = {'date': end.strftime('%Y%m%d')}
            manifest[period] = {'start': start.isoformat(), 'end': end.isoformat()}
            fetch('dateAvgRecentPrice.do', f'{period}/national.csv', params, dates)
            fetch('datePollAvgRecentPrice.do', f'{period}/brands.csv', params, dates)
            for code in codes:
                fetch('dateAreaAvgRecentPrice.do', f'{period}/areas/{code}.csv',
                      {**params, 'area': code}, dates)
        manifest['status'] = 'complete'
    except Exception as exc:
        manifest['status'] = 'failed'
        manifest['error'] = str(exc).replace(key, '[REDACTED]')
        print(manifest['error'], flush=True)
    finally:
        save_json(output / 'manifest.json', manifest)
    if manifest['status'] != 'complete':
        raise RuntimeError('데이터 수집 실패. 엑셀 생성과 메일 발송을 중단합니다.')
    print(f"Saved {len(manifest['files'])} data files; {len(manifest['warnings'])} coverage warnings.")
    return output


def read_data(folder, as_of):
    datasets = defaultdict(list)
    seen = set()
    csv_folder = folder / 'csv' if (folder / 'csv').is_dir() else folder
    manifest_path = folder / 'manifest.json'
    paths = None
    if manifest_path.is_file():
        record = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
        entries = record.get('files', [])
        if record.get('status') != 'complete':
            raise ValueError('수집 기록이 완료 상태가 아닙니다. 불완전한 데이터로 진행하지 않습니다.')
        if entries and all(entry['path'].endswith('.csv') for entry in entries):
            paths = [folder / entry['path'] for entry in entries]
            if any(not p.resolve().is_relative_to(folder.resolve()) for p in paths):
                raise ValueError('수집 기록에 데이터 폴더 외부 경로가 있습니다.')
    # For the previously converted CSV set, the old manifest lists JSON files.
    for filename in paths if paths is not None else sorted(csv_folder.rglob('*.csv')):
        with filename.open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                if 'DATE' not in row:  # Region-code response is a lookup, not a price record.
                    continue
                day = datetime.strptime(row['DATE'], '%Y%m%d').date()
                age = (as_of - day).days
                if not 1 <= age <= 14:
                    raise ValueError(f'Date outside comparison periods: {filename}, {day}')
                period = '최근 7일' if age <= 7 else '직전 7일'
                product = row['PRODCD']
                kind = '지역별 일별' if 'AREA_CD' in row else '상표별 일별' if 'POLL_DIV_CD' in row else '전국 일별'
                category = row.get('AREA_CD', row.get('POLL_DIV_CD', ''))
                key = kind, day, product, category
                if key in seen:
                    raise ValueError(f'Duplicate source record: {key}')
                seen.add(key)
                values = [period, day, PRODUCTS.get(product, product), float(row['PRICE']), product]
                if kind == '지역별 일별':
                    values += [row['AREA_NM'], category]
                elif kind == '상표별 일별':
                    values += [BRANDS.get(category, category), category]
                datasets[kind].append(values)
    for kind in ['전국 일별', '상표별 일별', '지역별 일별']:
        if not datasets[kind]:
            raise ValueError(f'Missing data: {kind}')
        datasets[kind].sort(key=lambda row: (row[1], row[4], row[6] if len(row) > 5 else ''))
    return datasets


def header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='맑은 고딕', size=10, color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def write_sources(wb, datasets):
    for index, kind in enumerate(['전국 일별', '상표별 일별', '지역별 일별'], 1):
        rows = datasets[kind]
        ws = wb.create_sheet(kind)
        ws.sheet_view.showGridLines = False
        ws['A1'] = '출처: 한국석유공사 오피넷 API / 저장된 CSV / 가격 단위: 원/L'
        labels = ['기간 구분', '날짜', '제품', '일평균 가격(원/L)', '제품코드']
        if kind == '지역별 일별':
            labels += ['지역', '지역코드']
        elif kind == '상표별 일별':
            labels += ['상표', '상표코드']
        header(ws, 3, labels)
        for values in rows:
            ws.append(values)
        for cells in ws.iter_rows(min_row=4):
            for cell in cells:
                cell.font = Font(name='맑은 고딕', size=10)
                cell.alignment = Alignment(vertical='center')
            cells[1].number_format = 'yyyy-mm-dd'
            cells[1].alignment = Alignment(horizontal='center', vertical='center')
            cells[3].number_format = PRICE_FORMAT
            cells[4].number_format = '@'
            if len(cells) > 5:
                cells[6].number_format = '@'
            ws.row_dimensions[cells[0].row].height = 20
        for col, width in {'A': 14, 'B': 15, 'C': 20, 'D': 23, 'E': 14, 'F': 24, 'G': 14}.items():
            ws.column_dimensions[col].width = width
        table = Table(displayName=f'PriceData{index}', ref=f'A3:{chr(64 + len(labels))}{ws.max_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = 'C4'


def make_summary(wb, datasets, as_of):
    ws = wb.active
    ws.title = '기간별 비교'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws['A2'] = '유가 기간별 비교'
    ws['A2'].font = Font(name='맑은 고딕', size=16, bold=True, color=NAVY)
    ws['A3'] = f"직전: {as_of - timedelta(days=14):%Y-%m-%d} ~ {as_of - timedelta(days=8):%Y-%m-%d}"
    ws['D3'] = f"최근: {as_of - timedelta(days=7):%Y-%m-%d} ~ {as_of - timedelta(days=1):%Y-%m-%d}"
    ws['A4'] = '7개 일평균 가격의 단순평균입니다. 공식 주간 통계와 다를 수 있습니다.'
    ws['A5'] = '관측일이 7일 미만이면 평균을 표시하지 않습니다. 0일은 해당 제품 자료가 없는 경우입니다.'
    ws['A7'] = '전국 제품별'
    labels = ['제품', '직전 평균(원/L)', '최근 평균(원/L)', '증감액(원/L)', '증감률', '직전 관측일', '최근 관측일']
    header(ws, 8, labels)
    ws['A15'] = '비교할 제품'
    ws['B15'] = '보통휘발유'
    ws['B15'].fill = PatternFill('solid', fgColor='FFF2CC')
    selector = DataValidation(type='list', formula1='"' + ','.join(PRODUCTS.values()) + '"')
    selector.errorTitle = '제품 선택'
    selector.error = '목록에 있는 제품을 선택하세요.'
    selector.showErrorMessage = True
    selector.errorStyle = 'stop'
    ws.add_data_validation(selector)
    selector.add(ws['B15'])

    def section(source, start, items, category=False):
        last = 3 + len(datasets[source])
        prefix = f"'{source}'!"
        prices = f'{prefix}$D$4:$D${last}'
        periods = f'{prefix}$A$4:$A${last}'
        products = f'{prefix}$C$4:$C${last}'
        categories = f'{prefix}$F$4:$F${last}'
        for r, item in enumerate(items, start):
            ws.cell(r, 1, item)
            product = '$B$15' if category else f'$A{r}'
            extra = f',{categories},$A{r}' if category else ''
            for avg_col, count_col, period in [('B', 'F', '직전 7일'), ('C', 'G', '최근 7일')]:
                criteria = f'{periods},"{period}",{products},{product}{extra}'
                ws[f'{count_col}{r}'] = f'=COUNTIFS({criteria})'
                ws[f'{avg_col}{r}'] = f'=IF({count_col}{r}=7,AVERAGEIFS({prices},{criteria}),"")'
                ws[f'{avg_col}{r}'].number_format = PRICE_FORMAT
            ws[f'D{r}'] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),C{r}-B{r},"")'
            ws[f'E{r}'] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(B{r})),IF(B{r}=0,"",D{r}/B{r}),"")'
            ws[f'D{r}'].number_format = '+#,##0.00;-#,##0.00;0.00'
            ws[f'E{r}'].number_format = '+0.00%;-0.00%;0.00%'
        end = start + len(items) - 1
        for operator, color in [('greaterThan', 'C00000'), ('lessThan', '0070C0')]:
            ws.conditional_formatting.add(f'D{start}:E{end}', CellIsRule(
                operator=operator, formula=['0'], font=Font(color=color)))

    section('전국 일별', 9, list(PRODUCTS.values()))
    brands = sorted({row[5] for row in datasets['상표별 일별']})
    ws['A17'] = '전국 제품·상표별'
    header(ws, 18, ['상표'] + labels[1:])
    section('상표별 일별', 19, brands, True)
    region_title = 19 + len(brands) + 2
    ws.cell(region_title, 1, '제품·지역별')
    header(ws, region_title + 1, ['지역'] + labels[1:])
    regions = sorted({row[5] for row in datasets['지역별 일별']})
    section('지역별 일별', region_title + 2, regions, True)
    ws.cell(ws.max_row + 2, 1, '전국 평균은 전국 데이터를 사용합니다. 지역별 평균을 다시 평균 내지 않습니다.')
    for row in ws.iter_rows():
        for cell in row:
            if cell.row != 2 and cell.fill.fgColor.rgb != '00203864':
                cell.font = Font(name='맑은 고딕', size=10)
            is_header = cell.row in [8, 18, region_title + 1]
            cell.alignment = Alignment(vertical='center', wrap_text=is_header,
                                       horizontal='center' if is_header else 'right' if cell.data_type == 'f' else 'left')
        if row[0].row not in [8, 18, region_title + 1]:
            ws.row_dimensions[row[0].row].height = 24
    for col, width in {'A': 26, 'B': 22, 'C': 22, 'D': 23, 'E': 17, 'F': 16, 'G': 16}.items():
        ws.column_dimensions[col].width = width
    ws.sheet_view.zoomScale = 85


def build_excel(folder, as_of):
    datasets = read_data(folder, as_of)
    workbook = Workbook()
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    write_sources(workbook, datasets)
    make_summary(workbook, datasets, as_of)
    destination = folder / '유가_기간비교.xlsx'
    temporary = destination.with_suffix('.tmp.xlsx')
    try:
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--to', required=True, nargs='+', help='받는 주소 (공백 또는 쉼표로 구분)')
    parser.add_argument('--as-of', type=date.fromisoformat, default=datetime.now(KST).date(),
                        help='수집기준일 YYYY-MM-DD (기본: 한국 시간 오늘)')
    parser.add_argument('--subject', help='메일 제목')
    parser.add_argument('--body', help='메일 본문')
    parser.add_argument('--skip-collect', action='store_true', help='API 호출 없이 해당 기준일의 기존 CSV 사용')
    parser.add_argument('--dry-run', action='store_true', help='엑셀 생성과 메일 구성까지 진행, SMTP 연결·전송 생략')
    args = parser.parse_args()
    stage = '설정 확인'
    try:
        recipients = list(dict.fromkeys(address.strip() for item in args.to
                                        for address in item.split(',') if address.strip()))
        if not recipients:
            raise ValueError('받는 이메일 주소를 입력하세요.')
        for recipient in recipients:
            validate_address(recipient)
        sender, password = load_settings()
        key = None if args.skip_collect else load_key()
        folder = ROOT / 'data' / 'opinet' / args.as_of.isoformat()
        stage = '데이터 수집'
        if args.skip_collect:
            print('[1/3] 기존 CSV 사용 (API 호출 없음)', flush=True)
            if not folder.is_dir():
                raise FileNotFoundError('지정한 기준일의 데이터 폴더가 없습니다.')
        else:
            print('[1/3] 오피넷 데이터 수집', flush=True)
            collect_data(args.as_of, key, folder)
        stage = '엑셀 생성'
        print('[2/3] 기간별 비교 엑셀 생성', flush=True)
        attachment = build_excel(folder, args.as_of)
        print(f'저장: {attachment}', flush=True)
        stage = '이메일 발송'
        message = build_message(sender, recipients, attachment, args.subject, args.body)
        if args.dry_run:
            print('[3/3] 메일 구성 확인 완료 (SMTP 연결·전송 없음)', flush=True)
            return 0
        print('[3/3] Gmail로 엑셀 발송', flush=True)
        with smtplib.SMTP_SSL('smtp.gmail.com', 465,
                             context=ssl.create_default_context(), timeout=30) as smtp:
            smtp.login(sender, password)
            rejected = smtp.send_message(message, from_addr=sender, to_addrs=recipients)
        if rejected:
            print('일부 수신자가 거부되었습니다. 나머지는 발송됐을 수 있으니 전체 재전송하지 마세요.', file=sys.stderr)
            return 1
        print('완료: Gmail SMTP 서버가 메일을 접수했습니다.', flush=True)
        return 0
    except smtplib.SMTPAuthenticationError:
        print('Gmail 인증 실패: email_address와 gmail_app_password를 확인하세요.', file=sys.stderr)
    except smtplib.SMTPRecipientsRefused:
        print('Gmail 서버가 받는 사람을 거부했습니다.', file=sys.stderr)
    except (TimeoutError, ConnectionError, smtplib.SMTPServerDisconnected):
        print(f'{stage} 중 연결이 중단되었습니다. 발송 단계였다면 보낸메일함을 확인하세요.', file=sys.stderr)
    except PermissionError:
        print(f'{stage} 실패: 엑셀이 열려 있다면 닫고, 파일 쓰기 권한을 확인하세요.', file=sys.stderr)
    except Exception as exc:
        # Do not log exception messages that might contain credentials or server replies.
        print(f'{stage} 실패 ({type(exc).__name__}). 다음 단계는 실행하지 않습니다.', file=sys.stderr)
    return 1


if __name__ == '__main__':
    raise SystemExit(main())
