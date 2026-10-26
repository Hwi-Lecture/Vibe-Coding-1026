"""Build a standalone HTML dashboard from a saved workbook; no API calls."""
import argparse
import json
import math
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def latest_workbook():
    candidates = []
    for path in (ROOT / 'data' / 'opinet').glob('*/유가_기간비교.xlsx'):
        try:
            candidates.append((date.fromisoformat(path.parent.name), path))
        except ValueError:
            continue
    if not candidates:
        raise FileNotFoundError('비교할 엑셀이 없습니다. --source로 파일 경로를 지정하세요.')
    return max(candidates)[1]


def load_records(source):
    workbook = load_workbook(source, read_only=True, data_only=True)
    records, seen = [], set()
    try:
        for name, kind in [('전국 일별', 'national'), ('상표별 일별', 'brand'), ('지역별 일별', 'region')]:
            if name not in workbook.sheetnames:
                raise ValueError(f'필수 시트가 없습니다: {name}')
            for values in workbook[name].iter_rows(min_row=4, values_only=True):
                if len(values) < 5 or values[1] is None:
                    continue
                period, day, product, price, code = values[:5]
                if not isinstance(day, (date, datetime)) or period not in {'최근 7일', '직전 7일'}:
                    raise ValueError(f'{name}: 날짜 또는 기간 구분을 확인하세요.')
                day = day.date() if isinstance(day, datetime) else day
                price = float(price)
                if not math.isfinite(price):
                    raise ValueError(f'{name}: 정상적인 가격이 아닙니다.')
                category = values[5] if len(values) > 5 else '전국'
                identity = (kind, day, code, category)
                if identity in seen:
                    raise ValueError(f'중복된 일별 자료입니다: {identity}')
                seen.add(identity)
                records.append({'kind': kind, 'period': 'recent' if period == '최근 7일' else 'previous',
                                'date': day.isoformat(), 'product': product, 'code': code,
                                'price': price, 'category': category})
    finally:
        workbook.close()
    if {r['kind'] for r in records} != {'national', 'brand', 'region'}:
        raise ValueError('전국·상표·지역별 자료가 모두 필요합니다.')
    return records


def build(source, output, template=None):
    records = load_records(source)
    last_day = date.fromisoformat(max(r['date'] for r in records))
    labels = {
        '__PREVIOUS_START__': (last_day - timedelta(days=13)).isoformat(),
        '__PREVIOUS_END__': (last_day - timedelta(days=7)).isoformat(),
        '__RECENT_START__': (last_day - timedelta(days=6)).isoformat(),
        '__RECENT_END__': last_day.isoformat(),
        '__AS_OF__': (last_day + timedelta(days=1)).isoformat(),
    }
    for period, start in [('previous', '__PREVIOUS_START__'), ('recent', '__RECENT_START__')]:
        expected = {(date.fromisoformat(labels[start]) + timedelta(days=i)).isoformat() for i in range(7)}
        actual = {r['date'] for r in records if r['period'] == period}
        if actual != expected:
            raise ValueError('두 기간은 연속된 두 개의 7일 구간이어야 합니다.')
    for marker in ['__PREVIOUS_START__', '__PREVIOUS_END__', '__RECENT_START__', '__RECENT_END__']:
        labels[marker[:-2] + '_SHORT__'] = labels[marker][5:].replace('-', '.')
    first_year = date.fromisoformat(labels['__PREVIOUS_START__']).year
    labels['__YEAR_LABEL__'] = (str(last_day.year) if first_year == last_day.year
                               else f'{first_year}~{last_day.year}') + '년'
    html = Path(template or ROOT / 'template.html').read_text(encoding='utf-8-sig')
    for marker, value in labels.items():
        html = html.replace(marker, value)
        html = html.replace(marker.replace('__', '__DOT_', 1), value.replace('-', '.'))
    html = html.replace('__TREND_PERIOD__', labels['__PREVIOUS_START__'].replace('-', '.') +
                        ' — ' + labels['__RECENT_END__'].replace('-', '.'))
    payload = json.dumps(records, ensure_ascii=False).replace('&', r'\u0026').replace('<', r'\u003c').replace('>', r'\u003e')
    html = html.replace('__DATA__', payload)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html, encoding='utf-8')
    return len(records)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, help='원본 엑셀 파일')
    parser.add_argument('--output', type=Path, default=ROOT / 'html_dashboard.html')
    args = parser.parse_args()
    count = build(args.source or latest_workbook(), args.output)
    print(f'완성 화면: {args.output.resolve()} / 일별 자료 {count}건')


if __name__ == '__main__':
    main()
