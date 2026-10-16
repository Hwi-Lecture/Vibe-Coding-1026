import io
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_COLUMNS = ["판매일자", "지점", "거래처명", "제품군", "모델명", "수량", "단가",
               "판매금액", "담당자", "연락처", "비고", "출처파일", "출처시트"]

COLUMN_ALIASES = {
    "판매일자": ["판매일자", "일자", "날짜"],
    "지점": ["지점", "지점명"],
    "거래처명": ["거래처명", "거래처"],
    "제품군": ["제품군", "제품종류"],
    "모델명": ["모델명", "모델"],
    "수량": ["수량", "판매수량"],
    "단가": ["단가", "판매단가"],
    "판매금액": ["판매금액", "금액"],
    "담당자": ["담당자", "영업담당"],
    "연락처": ["연락처"],
    "비고": ["비고", "특이사항"],
}
ALIAS_TO_STD = {a: std for std, aliases in COLUMN_ALIASES.items() for a in aliases}
BRANCH_ORDER = ["서울강남", "서울강북", "수원", "대전"]
PRODUCT_ORDER = ["가스 콘덴싱 보일러", "가스 보일러", "기름 보일러", "전기 보일러"]


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    m = re.search(r"-?\d[\d,]*", str(value))
    if not m:
        return None
    return int(m.group().replace(",", ""))


def to_date(value):
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", str(value))
    if not m:
        return None
    return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def clean_text(value):
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def clean_phone(value):
    text = clean_text(value)
    if not text:
        return None
    digits = re.sub(r"\D", "", text)
    if digits.startswith("02"):
        head, rest = digits[:2], digits[2:]
    else:
        head, rest = digits[:3], digits[3:]
    if len(rest) < 7:
        return text
    return f"{head}-{rest[:-4]}-{rest[-4:]}"


def branch_from_filename(filename):
    stem = re.sub(r"\.xlsx$", "", filename, flags=re.I)
    return re.sub(r"지점$", "", stem.split("_")[0]).strip()


def normalize_branch(value, filename):
    text = clean_text(value)
    if not text:
        return branch_from_filename(filename)
    text = text.replace(" ", "")
    return re.sub(r"지점$", "", text)


def find_header(ws):
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        mapped = {}
        for col, cell in enumerate(row):
            std = ALIAS_TO_STD.get(clean_text(cell) or "")
            if std and std not in mapped.values():
                mapped[col] = std
        if {"판매일자", "거래처명", "수량", "판매금액"} <= set(mapped.values()):
            return idx, mapped
    return None, None


def load_records(files):
    records, report = [], []
    for filename, data in files:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            header_row, mapped = find_header(ws)
            if header_row is None:
                continue
            for excel_row, row in enumerate(
                ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
            ):
                raw = {std: row[col] if col < len(row) else None for col, std in mapped.items()}
                date = to_date(raw.get("판매일자"))
                if date is None:
                    continue
                rec = {
                    "판매일자": date,
                    "지점": normalize_branch(raw.get("지점"), filename),
                    "거래처명": clean_text(raw.get("거래처명")),
                    "제품군": clean_text(raw.get("제품군")),
                    "모델명": clean_text(raw.get("모델명")),
                    "수량": to_number(raw.get("수량")),
                    "단가": to_number(raw.get("단가")),
                    "판매금액": to_number(raw.get("판매금액")),
                    "담당자": clean_text(raw.get("담당자")),
                    "연락처": clean_phone(raw.get("연락처")),
                    "비고": clean_text(raw.get("비고")),
                    "출처파일": filename,
                    "출처시트": ws.title,
                }
                label = f"{date.month}월 / {rec['거래처명'] or '(거래처 없음)'} / {rec['모델명']} / 수량 {clean_text(raw.get('수량'))}"

                def log(action, reason):
                    report.append({"출처파일": filename, "출처시트": ws.title, "엑셀 행": excel_row,
                                   "처리": action, "사유": reason, "내용": label})

                if not rec["거래처명"]:
                    log("제외", "거래처명 누락")
                    continue
                if not isinstance(raw.get("수량"), (int, float)) and rec["수량"] is None:
                    log("제외", "수량이 숫자가 아님")
                    continue
                if rec["수량"] is not None and rec["단가"] is not None:
                    calc = rec["수량"] * rec["단가"]
                    if rec["판매금액"] != calc:
                        before = rec["판매금액"]
                        rec["판매금액"] = calc
                        log("금액 수정", f"원본 금액 {before:,} → 계산값 {calc:,}" if before is not None
                            else f"원본 금액 없음 → 계산값 {calc:,}")
                records.append(rec)
    return records, report


def branch_key(name):
    return BRANCH_ORDER.index(name) if name in BRANCH_ORDER else len(BRANCH_ORDER)


def summarize(records):
    by_branch_month = {}
    by_product = {}
    for r in records:
        b = by_branch_month.setdefault(r["지점"], {})
        m = b.setdefault(r["판매일자"].month, [0, 0, 0])
        m[0] += 1
        m[1] += r["수량"]
        m[2] += r["판매금액"]
        p = by_product.setdefault(r["제품군"], [0, 0, 0])
        p[0] += 1
        p[1] += r["수량"]
        p[2] += r["판매금액"]
    return by_branch_month, by_product


def process(files):
    records, report = load_records(files)
    records.sort(key=lambda r: (r["판매일자"], r["출처파일"]))
    by_branch_month, by_product = summarize(records)
    branches = sorted(by_branch_month, key=branch_key)
    months = sorted({m for b in by_branch_month.values() for m in b})
    products = sorted(by_product, key=lambda p: PRODUCT_ORDER.index(p) if p in PRODUCT_ORDER else 99)
    branch_rows = []
    for b in branches:
        row = {"지점": b}
        total = [0, 0, 0]
        for m in months:
            c, q, a = by_branch_month[b].get(m, [0, 0, 0])
            row[f"{m}월 건수"], row[f"{m}월 수량"], row[f"{m}월 판매금액"] = c, q, a
            total = [total[0] + c, total[1] + q, total[2] + a]
        row["합계 건수"], row["합계 수량"], row["합계 판매금액"] = total
        branch_rows.append(row)
    product_rows = [{"제품군": p, "건수": by_product[p][0], "수량": by_product[p][1],
                     "판매금액": by_product[p][2]} for p in products]
    stats = {
        "원본 데이터 행": len(records) + sum(1 for r in report if r["처리"] == "제외"),
        "제외한 행": sum(1 for r in report if r["처리"] == "제외"),
        "금액 수정 행": sum(1 for r in report if r["처리"] == "금액 수정"),
        "정리결과 행": len(records),
        "총 수량": sum(r["수량"] for r in records),
        "총 판매금액": sum(r["판매금액"] for r in records),
    }
    return {"records": records, "report": report, "branch_rows": branch_rows,
            "product_rows": product_rows, "stats": stats, "months": months}


def _write_sheet(wb, title, headers, rows, number_cols=(), date_cols=(), total_row=False):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if total_row and rows:
        totals = ["합계"] + [sum(r[i] for r in rows) for i in range(1, len(headers))]
        ws.append(totals)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        width = max(len(str(header)) * 2, 12)
        for cell in ws[letter][1:60]:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) * 1.6 + 2, 48))
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            if col_idx in number_cols:
                cell.number_format = "#,##0"
            if col_idx in date_cols:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    return ws


def build_workbook(result):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rows = [[r[c] for c in OUT_COLUMNS] for r in result["records"]]
    _write_sheet(wb, "정리결과", OUT_COLUMNS, rows, number_cols=(6, 7, 8), date_cols=(1,))
    rep_headers = ["출처파일", "출처시트", "엑셀 행", "처리", "사유", "내용"]
    _write_sheet(wb, "검증리포트", rep_headers, [[r[h] for h in rep_headers] for r in result["report"]])
    months = result["months"]
    bm_headers = ["지점"] + [f"{m}월 {k}" for m in months for k in ("건수", "수량", "판매금액")] \
        + ["합계 건수", "합계 수량", "합계 판매금액"]
    bm_rows = [[r[h] for h in bm_headers] for r in result["branch_rows"]]
    _write_sheet(wb, "지점별_월별", bm_headers, bm_rows, number_cols=range(2, len(bm_headers) + 1), total_row=True)
    p_headers = ["제품군", "건수", "수량", "판매금액"]
    _write_sheet(wb, "제품군별", p_headers, [[r[h] for h in p_headers] for r in result["product_rows"]],
                 number_cols=(2, 3, 4), total_row=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
